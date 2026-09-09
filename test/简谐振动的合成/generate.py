# -*- coding: utf-8 -*-
"""简谐振动的合成 — 数据处理脚本。

注意：本实验公众号文章因微信限流无法抓取，示例数据基于教材原理构造，
  待获取公众号真实数据后可替换验证。按用户规则暂放 test 文件夹。
实验原理：两个相互垂直的简谐振动合成李萨如图形，
  频率比 f_y/f_x = N_x/N_y（N_x 为水平切点数，N_y 为垂直切点数）。
"""

import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(SCRIPT_DIR))
from common import *
from common.docx_report import DocxReportWriter

# 示例数据（基于教材原理构造，非公众号数据）
FX_DEFAULT = 100.0  # 标准频率 f_x，Hz
# 每组：设定频率比描述, N_x(水平切点数), N_y(垂直切点数)
DATA_DEFAULT = [
    ("1:1", 1, 1),
    ("2:1", 2, 1),
    ("3:2", 3, 2),
    ("4:3", 4, 3),
    ("1:2", 1, 2),
]


def _create_template(excel_path: str):
    if os.path.exists(excel_path):
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    title_font = Font(name="宋体", size=12, bold=True)
    header_font = Font(name="宋体", size=11, bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "李萨如图形"
    ws.merge_cells("A1:F1")
    ws["A1"] = "简谐振动的合成  数据记录"
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws["A3"] = "标准频率 f_x / Hz"
    ws["B3"] = FX_DEFAULT
    ws["A3"].font = header_font
    ws["A3"].border = border
    ws["B3"].border = border
    ws["B3"].alignment = center

    headers = ["设定频率比", "N_x (水平切点)", "N_y (垂直切点)",
               "计算频率比 f_y/f_x", "f_x / Hz", "f_y 计算值 / Hz"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=j, value=h)
        c.font = header_font; c.border = border; c.alignment = center

    for i, (ratio, nx, ny) in enumerate(DATA_DEFAULT):
        row = 6 + i
        ws.cell(row=row, column=1, value=ratio).border = border
        ws.cell(row=row, column=2, value=nx).border = border
        ws.cell(row=row, column=3, value=ny).border = border
        ws.cell(row=row, column=4).border = border  # 留空程序算
        ws.cell(row=row, column=5).border = border  # 留空程序算
        ws.cell(row=row, column=6).border = border  # 留空程序算
        for j in range(1, 7):
            ws.cell(row=row, column=j).alignment = center

    for col, w in zip("ABCDEF", [16, 18, 18, 20, 12, 16]):
        ws.column_dimensions[col].width = w

    wb.save(excel_path)


def _compute(excel_path: str) -> dict:
    f_x = float(read_cell(excel_path, sheet="李萨如图形", cell="B3"))

    ratios = []
    n_x_list = []
    n_y_list = []
    for row in range(6, 11):
        rv = read_cell(excel_path, sheet="李萨如图形", cell=f"A{row}")
        nx = read_cell(excel_path, sheet="李萨如图形", cell=f"B{row}")
        ny = read_cell(excel_path, sheet="李萨如图形", cell=f"C{row}")
        if rv is not None:
            ratios.append(str(rv))
            n_x_list.append(int(nx))
            n_y_list.append(int(ny))

    n = len(ratios)
    calc_ratios = [n_x_list[i] / n_y_list[i] for i in range(n)]
    f_y_list = [f_x * calc_ratios[i] for i in range(n)]

    return {
        "f_x": f_x, "ratios": ratios,
        "n_x": n_x_list, "n_y": n_y_list,
        "calc_ratios": calc_ratios, "f_y": f_y_list,
        "n": n,
    }


def _print_results(r: dict):
    print("=" * 60)
    print("简谐振动的合成 — 计算结果")
    print("=" * 60)
    print(f"标准频率 f_x = {r['f_x']} Hz")
    print(f"{'设定比':>8} {'N_x':>6} {'N_y':>6} {'计算比':>10} {'f_y(Hz)':>10}")
    for i in range(r["n"]):
        print(f"{r['ratios'][i]:>8} {r['n_x'][i]:6d} {r['n_y'][i]:6d} "
              f"{r['calc_ratios'][i]:10.4f} {r['f_y'][i]:10.2f}")
    print("=" * 60)


def _generate_docx(excel_path: str, output_path: str):
    r = _compute(excel_path)
    _print_results(r)

    doc = DocxReportWriter(output_path)
    doc.add_title("简谐振动的合成")
    doc.add_student_info()

    doc.add_heading("一、原始数据提交（拍照上传）", level=1)
    doc.add_paragraph("请在下方粘贴李萨如图形照片和原始数据记录。")

    doc.add_heading("二、数据处理", level=1)

    doc.add_heading("1. 实验原理", level=2)
    doc.add_paragraph("")
    doc.add_run("两个相互垂直的简谐振动：")
    doc.add_math(
        r"x = A_1 \cos(\omega_1 t + \varphi_1), \quad y = A_2 \cos(\omega_2 t + \varphi_2)"
    )
    doc.add_paragraph("")
    doc.add_run("当频率比为简单整数比时，合成轨迹为李萨如图形。频率比与图形切点数的关系：")
    doc.add_math(
        r"\frac{f_y}{f_x} = \frac{N_x}{N_y}"
    )
    doc.add_paragraph("")
    doc.add_run("其中 N_x 为图形与水平边界的切点数，N_y 为与垂直边界的切点数。")

    doc.add_heading("2. 测量数据与频率计算", level=2)
    doc.add_paragraph("")
    doc.add_run("标准频率 f_x = ")
    doc.add_inline_math(f"{r['f_x']:.0f} Hz")
    doc.add_run("。观察不同频率比下的李萨如图形，记录切点数并计算 f_y：")

    rows = []
    for i in range(r["n"]):
        rows.append([
            r["ratios"][i],
            str(r["n_x"][i]),
            str(r["n_y"][i]),
            f"{r['calc_ratios'][i]:.4f}",
            f"{r['f_x']:.0f}",
            f"{r['f_y'][i]:.2f}",
        ])
    doc.add_table(
        ["设定频率比", "N_x", "N_y", "计算 f_y/f_x", "f_x / Hz", "f_y / Hz"],
        rows, col_widths=[2.0, 1.5, 1.5, 2.2, 1.8, 2.0]
    )

    doc.add_heading("三、实验结果分析", level=1)
    doc.add_paragraph("")
    doc.add_run("本实验通过观察李萨如图形测量了不同频率比下的未知频率。利用切点数法计算频率比，")
    doc.add_run("方法简便直观，不需要测量相位差。实验结果表明，当频率比为简单整数比时，")
    doc.add_run("李萨如图形稳定且闭合，切点数法可以准确测定频率比。")

    doc.add_paragraph("")
    doc.add_run("误差来源分析：")
    doc.add_run("（1）李萨如图形的切点判断存在视差，特别是图形较复杂时切点数不易数清；")
    doc.add_run("（2）两个信号源的频率稳定性影响图形的稳定度，频率漂移会导致图形缓慢滚动；")
    doc.add_run("（3）示波器的水平和垂直增益不对称会导致图形畸变，影响切点判断；")
    doc.add_run("（4）相位差的变化会改变图形形状，但不影响切点数和频率比。")

    doc.add_heading("四、思考题", level=1)

    doc.add_heading("1. 为什么李萨如图形可以用来测量频率？", level=2)
    doc.add_paragraph(
        "答：当两个相互垂直的简谐振动频率比为简单整数比时，合成轨迹是稳定的闭合曲线（李萨如图形）。"
        "图形在水平方向的切点数 N_x 对应 y 方向振动的频率，垂直方向的切点数 N_y 对应 x 方向振动的频率，"
        "因此 f_y/f_x = N_x/N_y。已知一个标准频率，通过数切点数就可以计算另一个频率，"
        "这种方法不需要测量时间和周期，直观且精度较高。"
    )

    doc.add_heading("2. 当两个频率完全相等但相位差不同时，李萨如图形如何变化？", level=2)
    doc.add_paragraph(
        "答：当 f_x = f_y 时，李萨如图形为椭圆（或退化为直线/圆），具体形状取决于相位差 Δφ："
        "Δφ = 0 或 π 时退化为直线（斜率为正或负）；Δφ = π/2 或 3π/2 时为正椭圆（振幅相等时为圆）；"
        "其他相位差时为倾斜椭圆。相位差只改变图形的形状和取向，不影响切点数（N_x = N_y = 1）。"
    )

    doc.save()
    doc.close()


def main():
    EXCEL_FILE = os.path.join(SCRIPT_DIR, "数据.xlsx")
    DOCX_FILE = os.path.join(SCRIPT_DIR, "简谐振动的合成.docx")
    if not os.path.exists(EXCEL_FILE):
        _create_template(EXCEL_FILE)
        print("已生成数据模板（示例数据），可修改后重新运行。")
    _generate_docx(EXCEL_FILE, DOCX_FILE)
    print(f"报告已生成: {DOCX_FILE}")


if __name__ == "__main__":
    main()
